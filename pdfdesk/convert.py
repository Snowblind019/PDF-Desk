"""Import other formats into PDF and export PDFs to other formats. All offline.

Import:  images, text/code, Markdown, HTML, CSV, XPS, EPUB, MOBI, FB2, CBZ, SVG, and Office
         files (Word, Excel, PowerPoint, OpenDocument, RTF) through LibreOffice when installed.
Export:  Word, Excel, PowerPoint, OpenDocument text, RTF, PNG, JPEG, TIFF, WebP, SVG, text,
         HTML and Markdown."""
from __future__ import annotations

import csv
import html as _html
import io
import logging
import os
import re
import tempfile
from collections import Counter
from pathlib import Path

import pymupdf as fitz

from pdfdesk import externals, fonts
from pdfdesk.pdfops import Cancelled, Progress, _tick, _unique_path

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".jfif", ".bmp", ".gif", ".tif", ".tiff", ".webp", ".jxr",
              ".jpx", ".jp2", ".pnm", ".pbm", ".pgm", ".ppm", ".pam", ".psd", ".tga", ".ico", ".heic", ".avif"}
MUPDF_DOC_EXTS = {".xps", ".oxps", ".epub", ".mobi", ".fb2", ".cbz", ".svg"}
TEXT_EXTS = {".txt", ".text", ".log", ".json", ".xml", ".yaml", ".yml", ".ini", ".conf", ".cfg", ".toml",
             ".py", ".sh", ".bash", ".zsh", ".ps1", ".bat", ".cmd", ".tf", ".tfvars", ".hcl", ".js", ".ts",
             ".jsx", ".tsx", ".css", ".c", ".h", ".cpp", ".hpp", ".cs", ".java", ".kt", ".go", ".rs", ".rb",
             ".php", ".pl", ".lua", ".sql", ".r", ".swift", ".srt", ".vtt", ".env", ".properties", ".nfo"}
MARKDOWN_EXTS = {".md", ".markdown", ".mdown"}
HTML_EXTS = {".html", ".htm", ".xhtml"}
CSV_EXTS = {".csv", ".tsv"}
OFFICE_EXTS = {".doc", ".docx", ".docm", ".dot", ".dotx", ".odt", ".ott", ".rtf", ".wpd", ".pages",
               ".xls", ".xlsx", ".xlsm", ".ods", ".ots", ".numbers",
               ".ppt", ".pptx", ".pps", ".ppsx", ".odp", ".otp", ".key", ".odg", ".vsd", ".vsdx", ".pub"}
MUPDF_OFFICE_FALLBACK = {".docx", ".xlsx", ".pptx"}

ALL_IMPORT_EXTS = IMAGE_EXTS | MUPDF_DOC_EXTS | TEXT_EXTS | MARKDOWN_EXTS | HTML_EXTS | CSV_EXTS | OFFICE_EXTS

if hasattr(fitz, "no_recommend_layout"):  # silence an advert for an optional add-on package
    fitz.no_recommend_layout()

PAPER = {"letter": fitz.paper_rect("letter"), "a4": fitz.paper_rect("a4"), "legal": fitz.paper_rect("legal")}


class ConversionError(RuntimeError):
    pass


def ext_of(path: str) -> str:
    return Path(path).suffix.lower()


def can_open_directly(path: str) -> bool:
    return ext_of(path) == ".pdf"


def can_import(path: str) -> bool:
    return ext_of(path) in ALL_IMPORT_EXTS


def import_filter() -> str:
    def pats(exts):
        return " ".join(f"*{e}" for e in sorted(exts))
    groups = [
        ("All supported files", {".pdf"} | ALL_IMPORT_EXTS),
        ("PDF files", {".pdf"}),
        ("Office documents", OFFICE_EXTS),
        ("Images", IMAGE_EXTS),
        ("Text and code", TEXT_EXTS | CSV_EXTS),
        ("Markdown and web pages", MARKDOWN_EXTS | HTML_EXTS),
        ("E-books and other documents", MUPDF_DOC_EXTS),
    ]
    return ";;".join(f"{name} ({pats(exts)})" for name, exts in groups) + ";;All files (*)"


# =========================================================================== import

def to_pdf_bytes(path: str, paper: str = "letter", image_page: str = "image", progress: Progress = None) -> bytes:
    """Convert any supported file into PDF bytes. image_page: 'image', 'letter' or 'a4'."""
    ext = ext_of(path)
    if ext == ".pdf":
        with open(path, "rb") as fh:
            return fh.read()
    if ext in IMAGE_EXTS:
        return images_to_pdf_bytes([path], page_size=image_page, margin=0 if image_page == "image" else 36)
    if ext in MUPDF_DOC_EXTS:
        return mupdf_doc_to_pdf_bytes(path)
    if ext in MARKDOWN_EXTS:
        return markdown_to_pdf_bytes(_read_text(path), base_dir=os.path.dirname(path), paper=paper)
    if ext in HTML_EXTS:
        return html_to_pdf_bytes(_read_text(path), base_dir=os.path.dirname(path), paper=paper)
    if ext in CSV_EXTS:
        # Always the built-in table layout: a spreadsheet program would run formulas in the file.
        return csv_to_pdf_bytes(path, paper=paper)
    if ext in TEXT_EXTS:
        return text_to_pdf_bytes(_read_text(path), title=Path(path).name, paper=paper)
    if ext in OFFICE_EXTS:
        return office_to_pdf_bytes(path)
    # Last try: maybe MuPDF can open it anyway.
    try:
        return mupdf_doc_to_pdf_bytes(path)
    except Exception:
        raise ConversionError(f"PDF Desk can't convert {Path(path).suffix or 'this'} files.")


def files_to_pdf_bytes(paths: list[str], paper: str = "letter", image_page: str = "image",
                       progress: Progress = None) -> bytes:
    """Convert several files of any supported type and combine them into one PDF."""
    out = fitz.open()
    for i, path in enumerate(paths):
        _tick(progress, i, len(paths), f"Adding {Path(path).name}")
        data = to_pdf_bytes(path, paper=paper, image_page=image_page)
        src = fitz.open("pdf", data)
        if src.needs_pass:
            raise ConversionError(f"{Path(path).name} is password protected. Open it and remove the password first.")
        out.insert_pdf(src)
    _tick(progress, len(paths), len(paths), "Saving")
    if out.page_count == 0:
        raise ConversionError("No pages were produced.")
    return out.tobytes(garbage=3, deflate=True)


def _read_text(path: str) -> str:
    raw = Path(path).read_bytes()
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16", errors="replace")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def _image_bytes_for_mupdf(path_or_bytes) -> bytes:
    """Return image bytes MuPDF understands, converting via Pillow when needed."""
    data = path_or_bytes if isinstance(path_or_bytes, (bytes, bytearray)) else Path(path_or_bytes).read_bytes()
    try:
        test = fitz.open(stream=data, filetype="img" if isinstance(path_or_bytes, (bytes, bytearray))
                         else ext_of(path_or_bytes).lstrip("."))
        if test.page_count:
            return bytes(data)
    except Exception:
        pass
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(data))
        if img.mode not in ("RGB", "RGBA", "L", "LA"):
            img = img.convert("RGBA" if "A" in img.getbands() else "RGB")
        buf = io.BytesIO()
        img.save(buf, "PNG")
        return buf.getvalue()
    except Exception as exc:
        raise ConversionError(f"Could not read the image ({exc}).")


def images_to_pdf_bytes(items: list, page_size: str = "image", margin: float = 0,
                        progress: Progress = None) -> bytes:
    """items: file paths or raw image bytes. page_size: 'image' (page matches the image),
    'letter' or 'a4' (image fitted on the page, orientation chosen automatically)."""
    out = fitz.open()
    for i, item in enumerate(items):
        _tick(progress, i, len(items), "Adding images")
        data = _image_bytes_for_mupdf(item)
        img_doc = fitz.open(stream=data, filetype="img")
        for frame in range(img_doc.page_count):  # multi-page TIFF / animated GIF frames
            if page_size == "image":
                pdf = fitz.open("pdf", img_doc.convert_to_pdf(frame, frame))
                out.insert_pdf(pdf)
            else:
                rect = img_doc[frame].rect
                paper = PAPER.get(page_size, PAPER["letter"])
                if rect.width > rect.height:
                    paper = fitz.Rect(0, 0, paper.height, paper.width)
                page = out.new_page(width=paper.width, height=paper.height)
                box = fitz.Rect(margin, margin, paper.width - margin, paper.height - margin)
                single = fitz.open("pdf", img_doc.convert_to_pdf(frame, frame))
                page.show_pdf_page(box, single, 0, keep_proportion=True)
    if out.page_count == 0:
        raise ConversionError("No images could be read.")
    return out.tobytes(garbage=3, deflate=True)


def mupdf_doc_to_pdf_bytes(path: str) -> bytes:
    src = fitz.open(path)
    if src.is_reflowable:
        a5 = fitz.paper_rect("a5")
        src.layout(width=a5.width, height=a5.height, fontsize=11)
    if src.is_pdf:
        return src.tobytes()
    return src.convert_to_pdf()


BASE_CSS = """
body { font-family: sans-serif; font-size: 11pt; line-height: 1.4; color: #1d1d1f; }
h1 { font-size: 20pt; margin: 0 0 8pt 0; } h2 { font-size: 16pt; margin: 14pt 0 6pt 0; }
h3 { font-size: 13pt; margin: 12pt 0 4pt 0; } h4, h5, h6 { font-size: 11pt; }
p { margin: 0 0 8pt 0; }
pre, code { font-family: monospace; font-size: 9.5pt; }
pre { background-color: #f3f3f5; padding: 6pt; margin: 0 0 8pt 0; white-space: pre-wrap; }
table { border-collapse: collapse; margin: 0 0 8pt 0; }
th, td { border: 1px solid #b0b0b8; padding: 3pt 5pt; text-align: left; vertical-align: top; }
th { background-color: #ececf0; font-weight: bold; }
blockquote { margin: 0 0 8pt 12pt; color: #555; }
img { max-width: 100%; }
"""


def html_to_pdf_bytes(html: str, css: str = "", base_dir: str | None = None,
                      paper: "str | fitz.Rect" = "letter", margin: float = 54) -> bytes:
    """Lay out HTML with MuPDF's built-in HTML engine (no browser, no network)."""
    page_rect = paper if isinstance(paper, fitz.Rect) else PAPER.get(paper, PAPER["letter"])
    where = page_rect + (margin, margin, -margin, -margin)
    archive = fitz.Archive(base_dir) if base_dir and os.path.isdir(base_dir) else None
    # MuPDF's HTML engine has no network access and only reads files from base_dir and its subfolders.
    story = fitz.Story(html=html, user_css=BASE_CSS + css, archive=archive)
    buf = io.BytesIO()
    writer = fitz.DocumentWriter(buf)
    more = True
    pages = 0
    while more:
        device = writer.begin_page(page_rect)
        more, _filled = story.place(where)
        story.draw(device)
        writer.end_page()
        pages += 1
        if pages > 5000:
            break
    writer.close()
    return buf.getvalue()


def text_to_pdf_bytes(text: str, title: str = "", paper: str = "letter", mono: bool = True) -> bytes:
    body = _html.escape(text.replace("\r\n", "\n").replace("\t", "    "))
    tag = "pre" if mono else "div"
    css = "pre { background-color: transparent; padding: 0; font-size: 9.5pt; }" \
          "div { white-space: pre-wrap; }"
    return html_to_pdf_bytes(f"<{tag}>{body}</{tag}>", css=css, paper=paper, margin=48)


def markdown_to_pdf_bytes(md_text: str, base_dir: str | None = None, paper: str = "letter") -> bytes:
    try:
        import markdown
        html = markdown.markdown(md_text, extensions=["tables", "fenced_code", "sane_lists", "nl2br"])
    except Exception:
        html = f"<pre>{_html.escape(md_text)}</pre>"
    return html_to_pdf_bytes(html, base_dir=base_dir, paper=paper)


def csv_to_pdf_bytes(path: str, paper: str = "letter") -> bytes:
    text = _read_text(path)
    delim = "\t" if ext_of(path) == ".tsv" else None
    if delim is None:
        try:
            delim = csv.Sniffer().sniff(text[:4096]).delimiter
        except Exception:
            delim = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    if not rows:
        return text_to_pdf_bytes(text, paper=paper)
    head = "".join(f"<th>{_html.escape(c)}</th>" for c in rows[0])
    body = "".join("<tr>" + "".join(f"<td>{_html.escape(c)}</td>" for c in r) + "</tr>" for r in rows[1:])
    widest = max(len(r) for r in rows)
    css = "body { font-size: %spt; }" % (9 if widest <= 6 else 7)
    use_paper = PAPER.get(paper, PAPER["letter"])
    if widest > 6:  # wide tables read better in landscape
        use_paper = fitz.Rect(0, 0, use_paper.height, use_paper.width)
    return html_to_pdf_bytes(f"<table><tr>{head}</tr>{body}</table>", css=css, paper=use_paper, margin=36)


def office_to_pdf_bytes(path: str) -> bytes:
    ext = ext_of(path)
    if externals.libreoffice_available():
        return externals.libreoffice_convert(path, "pdf")
    if ext in MUPDF_OFFICE_FALLBACK:
        # MuPDF can read .docx/.xlsx/.pptx on its own, with simpler layout.
        return mupdf_doc_to_pdf_bytes(path)
    raise externals.ExternalToolMissing(
        f"Converting {ext} files needs LibreOffice, which was not found.\n\n"
        "Install LibreOffice (free) and try again. On Fedora: sudo dnf install libreoffice\n"
        "On Windows: install it from libreoffice.org, then restart PDF Desk.")


# =========================================================================== export

EXPORT_FORMATS = [
    # key, label, extension, kind ('file' = one output file, 'folder' = one file per page)
    ("docx", "Word document (.docx)", ".docx", "file"),
    ("xlsx", "Excel workbook (.xlsx)", ".xlsx", "file"),
    ("pptx", "PowerPoint presentation (.pptx)", ".pptx", "file"),
    ("odt", "OpenDocument text (.odt)", ".odt", "file"),
    ("rtf", "Rich Text Format (.rtf)", ".rtf", "file"),
    ("png", "PNG images", ".png", "folder"),
    ("jpg", "JPEG images", ".jpg", "folder"),
    ("webp", "WebP images", ".webp", "folder"),
    ("tiff", "TIFF image (all pages in one file)", ".tiff", "file"),
    ("svg", "SVG images", ".svg", "folder"),
    ("txt", "Plain text (.txt)", ".txt", "file"),
    ("html", "Web page (.html)", ".html", "file"),
    ("md", "Markdown (.md)", ".md", "file"),
]
NEEDS_LIBREOFFICE = {"odt", "rtf"}


def export_format(key: str):
    for fmt in EXPORT_FORMATS:
        if fmt[0] == key:
            return fmt
    raise KeyError(key)


def export_pdf(pdf_bytes: bytes, fmt: str, target: str, pages: list[int] | None = None,
               dpi: int = 150, quality: int = 90, page_markers: bool = False,
               progress: Progress = None) -> list[str]:
    """Export unencrypted PDF bytes. `target` is a file path, or a folder for per-page formats.
    Returns the list of files written."""
    doc = fitz.open("pdf", pdf_bytes)
    pages = list(range(doc.page_count)) if not pages else pages
    stem = Path(target).stem
    if fmt in ("png", "jpg", "webp", "svg"):
        os.makedirs(target, exist_ok=True)
        return _export_pages_as_images(doc, fmt, target, stem, pages, dpi, quality, progress)
    if fmt == "tiff":
        return [_export_tiff(doc, target, pages, dpi, progress)]
    if fmt == "txt":
        return [_export_text(doc, target, pages, page_markers, progress)]
    if fmt == "html":
        return [_export_html(doc, target, pages, progress)]
    if fmt == "md":
        return [_export_markdown(doc, target, pages, progress)]
    if fmt == "docx":
        return [_export_docx(pdf_bytes, target, pages, progress)]
    if fmt == "xlsx":
        return [_export_xlsx(doc, target, pages, progress)]
    if fmt == "pptx":
        return [_export_pptx(doc, target, pages, dpi, progress)]
    if fmt in NEEDS_LIBREOFFICE:
        return [_export_via_libreoffice(pdf_bytes, fmt, target, pages, progress)]
    raise ConversionError(f"Unknown export format {fmt}")


def _export_pages_as_images(doc, fmt, folder, stem, pages, dpi, quality, progress) -> list[str]:
    written = []
    digits = max(3, len(str(doc.page_count)))
    for i, pno in enumerate(pages):
        _tick(progress, i, len(pages), f"Page {pno + 1}")
        page = doc[pno]
        path = os.path.join(folder, f"{stem} - page {pno + 1:0{digits}d}.{fmt}")
        if fmt == "svg":
            with open(path, "w", encoding="utf-8") as fh:
                fh.write(page.get_svg_image(text_as_path=False))
        else:
            z = fonts.safe_zoom(page.rect, dpi / 72)
            pix = page.get_pixmap(matrix=fitz.Matrix(z, z), alpha=False)
            if fmt == "png":
                pix.save(path)
            elif fmt == "jpg":
                pix.save(path, jpg_quality=quality)
            else:
                pix.pil_save(path, format="WEBP", quality=quality)
        written.append(path)
    _tick(progress, len(pages), len(pages), "Done")
    return written


def _export_tiff(doc, target, pages, dpi, progress) -> str:
    from PIL import Image
    frames = []
    for i, pno in enumerate(pages):
        _tick(progress, i, len(pages), f"Page {pno + 1}")
        z = fonts.safe_zoom(doc[pno].rect, dpi / 72)
        pix = doc[pno].get_pixmap(matrix=fitz.Matrix(z, z), alpha=False)
        frames.append(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
    frames[0].save(target, save_all=True, append_images=frames[1:], compression="tiff_deflate",
                   dpi=(dpi, dpi))
    return target


def _export_text(doc, target, pages, page_markers, progress) -> str:
    parts = []
    for i, pno in enumerate(pages):
        _tick(progress, i, len(pages), f"Page {pno + 1}")
        text = doc[pno].get_text("text", sort=True, flags=fitz.TEXT_PRESERVE_WHITESPACE | fitz.TEXT_MEDIABOX_CLIP).rstrip()
        parts.append(f"===== Page {pno + 1} =====\n{text}" if page_markers else text)
    with open(target, "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n\n".join(parts) + "\n")
    return target


def _export_html(doc, target, pages, progress) -> str:
    title = _html.escape(doc.metadata.get("title") or Path(target).stem)
    out = [f"<!DOCTYPE html>\n<html><head><meta charset='utf-8'><title>{title}</title><style>"
           "body{background:#e9e9ec;margin:0;padding:24px;font-family:sans-serif}"
           ".page{background:#fff;max-width:900px;margin:0 auto 24px;padding:40px;"
           "box-shadow:0 1px 4px rgba(0,0,0,.2)}.page img{max-width:100%;height:auto}"
           ".pno{color:#888;font-size:12px;text-align:right}"
           "</style></head><body>"]
    for i, pno in enumerate(pages):
        _tick(progress, i, len(pages), f"Page {pno + 1}")
        body = doc[pno].get_text("xhtml")
        out.append(f"<section class='page' id='page-{pno + 1}'><div class='pno'>Page {pno + 1}</div>{body}</section>")
    out.append("</body></html>")
    with open(target, "w", encoding="utf-8") as fh:
        fh.write("\n".join(out))
    return target


# ---- Markdown ------------------------------------------------------------

_BULLETS = ("•", "◦", "▪", "●", "○", "■", "□", "–", "-", "*", "·", "")


def _body_size(doc, pages) -> float:
    sizes = Counter()
    for pno in pages[:30]:
        for block in doc[pno].get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for span in line["spans"]:
                    sizes[round(span["size"], 1)] += len(span["text"].strip())
    return sizes.most_common(1)[0][0] if sizes else 11.0


def _md_escape(text: str) -> str:
    """Escape characters that Markdown viewers would treat as formatting or raw HTML."""
    return text.replace("\\", "\\\\").replace("*", "\\*").replace("<", "\\<").replace(">", "\\>")


def _span_md(span) -> str:
    text = span["text"]
    if not text.strip():
        return text
    name = span["font"].lower()
    bold = span["flags"] & 16 or "bold" in name
    italic = span["flags"] & 2 or "italic" in name or "oblique" in name
    lead = text[: len(text) - len(text.lstrip())]
    trail = text[len(text.rstrip()):]
    core = _md_escape(text.strip())
    if bold and italic:
        core = f"***{core}***"
    elif bold:
        core = f"**{core}**"
    elif italic:
        core = f"*{core}*"
    return lead + core + trail


def _is_mono(spans) -> bool:
    return all((s["flags"] & 8) or "mono" in s["font"].lower() or "courier" in s["font"].lower()
               for s in spans if s["text"].strip())


def page_to_markdown(page: fitz.Page, body: float) -> str:
    items: list[tuple[float, float, str]] = []
    table_boxes = []
    try:
        tables = page.find_tables()
        for t in tables.tables:
            try:
                md = t.to_markdown(clean=False)
            except Exception:
                continue
            if md.strip():
                table_boxes.append(fitz.Rect(t.bbox))
                items.append((t.bbox[1], t.bbox[0], md.strip()))
    except Exception:
        pass
    for block in page.get_text("dict", sort=True, flags=fitz.TEXT_PRESERVE_WHITESPACE | fitz.TEXT_MEDIABOX_CLIP)["blocks"]:
        if block.get("type") != 0:
            continue
        bbox = fitz.Rect(block["bbox"])
        if any(tb.intersects(bbox) and (tb & bbox).get_area() > 0.5 * bbox.get_area() for tb in table_boxes):
            continue
        lines = [ln for ln in block.get("lines", []) if "".join(s["text"] for s in ln["spans"]).strip()]
        if not lines:
            continue
        all_spans = [s for ln in lines for s in ln["spans"]]
        if _is_mono(all_spans) and len(lines) > 1:
            code = "\n".join("".join(s["text"] for s in ln["spans"]).rstrip() for ln in lines)
            items.append((bbox.y0, bbox.x0, f"```\n{code}\n```"))
            continue
        max_size = max(s["size"] for s in all_spans if s["text"].strip())
        plain = " ".join("".join(s["text"] for s in ln["spans"]).strip() for ln in lines)
        if max_size >= body * 1.15 and len(lines) <= 3 and len(plain) < 200:
            ratio = max_size / body
            level = 1 if ratio >= 1.8 else 2 if ratio >= 1.4 else 3
            items.append((bbox.y0, bbox.x0, "#" * level + " " + _md_escape(plain)))
            continue
        out_lines: list[str] = []
        para = ""
        for ln in lines:
            text = "".join(_span_md(s) for s in ln["spans"]).strip()
            raw = "".join(s["text"] for s in ln["spans"]).strip()
            bullet = raw[:1] in _BULLETS and len(raw) > 1 and raw[1:2] in (" ", "\t", "")
            numbered = re.match(r"^\(?\d{1,3}[.)]\s", raw)
            if bullet or numbered:
                if para:
                    out_lines.append(para)
                para = ("- " + text.lstrip("".join(_BULLETS)).strip()) if bullet else text
                continue
            if para.endswith("-") and text[:1].islower():
                para = para[:-1] + text
            elif para:
                para += " " + text
            else:
                para = text
        if para:
            out_lines.append(para)
        items.append((bbox.y0, bbox.x0, "\n".join(out_lines)))
    items.sort(key=lambda it: (round(it[0]), it[1]))
    return "\n\n".join(it[2] for it in items)


def _export_markdown(doc, target, pages, progress) -> str:
    body = _body_size(doc, pages)
    parts = []
    for i, pno in enumerate(pages):
        _tick(progress, i, len(pages), f"Page {pno + 1}")
        parts.append(page_to_markdown(doc[pno], body))
    with open(target, "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n\n---\n\n".join(p for p in parts if p.strip()) + "\n")
    return target


# ---- Office --------------------------------------------------------------

def _export_docx(pdf_bytes, target, pages, progress) -> str:
    _tick(progress, 0, 2, "Analyzing layout (this can take a while on long files)")
    tmpdir = tempfile.mkdtemp(prefix="pdfdesk-")
    src = os.path.join(tmpdir, "input.pdf")
    with open(src, "wb") as fh:
        fh.write(pdf_bytes)
    try:
        from pdf2docx import Converter
        logging.getLogger().setLevel(logging.WARNING)
        cv = Converter(src)
        try:
            cv.convert(target, pages=pages)
        finally:
            cv.close()
    except ImportError:
        _basic_docx(fitz.open("pdf", pdf_bytes), target, pages)
    finally:
        try:
            os.remove(src)
            os.rmdir(tmpdir)
        except OSError:
            pass
    _tick(progress, 2, 2, "Done")
    return target


def _basic_docx(doc, target, pages) -> None:
    import docx
    out = docx.Document()
    for n, pno in enumerate(pages):
        for block in doc[pno].get_text("blocks", sort=True):
            if block[6] == 0 and block[4].strip():
                out.add_paragraph(block[4].strip().replace("\n", " "))
        if n < len(pages) - 1:
            out.add_page_break()
    out.save(target)


def _rows_from_text(page: fitz.Page) -> list[list[str]]:
    """Fallback for pages without detectable tables: one row per text line,
    with cells split where there are wide gaps between words."""
    words = page.get_text("words", sort=True)
    lines: dict[tuple, list] = {}
    for w in words:
        lines.setdefault((w[5], w[6]), []).append(w)
    rows = []
    for key in sorted(lines, key=lambda k: (lines[k][0][1], lines[k][0][0])):
        ws = sorted(lines[key], key=lambda w: w[0])
        cells, cur, last_x1 = [], [], None
        for w in ws:
            gap_limit = (w[3] - w[1]) * 1.2
            if last_x1 is not None and w[0] - last_x1 > gap_limit:
                cells.append(" ".join(cur))
                cur = []
            cur.append(w[4])
            last_x1 = w[2]
        if cur:
            cells.append(" ".join(cur))
        rows.append(cells)
    return rows


def _cell_value(text):
    if text is None:
        return ""
    s = str(text).strip()
    num = s.replace(",", "")
    if re.fullmatch(r"-?\d+", num) and len(num) < 16 and not (len(num) > 1 and num.startswith("0")):
        return int(num)
    if re.fullmatch(r"-?\d+\.\d+", num):
        try:
            return float(num)
        except ValueError:
            pass
    return s


def _export_xlsx(doc, target, pages, progress) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    found_tables = False
    sheets = []
    for i, pno in enumerate(pages):
        _tick(progress, i, len(pages), f"Looking for tables on page {pno + 1}")
        try:
            tables = doc[pno].find_tables().tables
        except Exception:
            tables = []
        for k, t in enumerate(tables):
            rows = t.extract()
            if rows and any(any(c for c in r) for r in rows):
                found_tables = True
                sheets.append((f"Page {pno + 1}" + (f" ({k + 1})" if len(tables) > 1 else ""), rows, True))
    if not found_tables:
        for pno in pages:
            rows = _rows_from_text(doc[pno])
            if rows:
                sheets.append((f"Page {pno + 1}", rows, False))
    if not sheets:
        sheets.append(("Page 1", [["(no text found - is this a scanned PDF? Run Recognize Text first)"]], False))
    used = set()
    for title, rows, header in sheets:
        name = title[:31]
        while name in used:
            name = (name[:28] + "_" + str(len(used)))[:31]
        used.add(name)
        ws = wb.create_sheet(name)
        widths: dict[int, int] = {}
        for r, row in enumerate(rows, start=1):
            for c, val in enumerate(row, start=1):
                text = (val or "").replace("\n", " ") if isinstance(val, str) or val is None else val
                cell = ws.cell(row=r, column=c, value=_cell_value(text))
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"  # text from the PDF must never become a live formula
                widths[c] = max(widths.get(c, 0), len(str(cell.value or "")))
                if header and r == 1:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill("solid", fgColor="ECECF0")
        for c, w in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = min(60, max(8, w + 2))
        if header:
            ws.freeze_panes = "A2"
    wb.save(target)
    return target


def _export_pptx(doc, target, pages, dpi, progress) -> str:
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    first = doc[pages[0]].rect
    max_pt = 56 * 72  # PowerPoint's largest slide side is 56 inches
    scale = min(1.0, max_pt / max(first.width, first.height))
    prs.slide_width = Emu(int(first.width * scale * 12700))
    prs.slide_height = Emu(int(first.height * scale * 12700))
    blank = prs.slide_layouts[6]
    for i, pno in enumerate(pages):
        _tick(progress, i, len(pages), f"Slide {i + 1}")
        page = doc[pno]
        z = fonts.safe_zoom(page.rect, dpi / 72)
        pix = page.get_pixmap(matrix=fitz.Matrix(z, z), alpha=False)
        slide = prs.slides.add_slide(blank)
        r = page.rect
        sw, sh = prs.slide_width, prs.slide_height
        ratio = min(sw / (r.width * 12700), sh / (r.height * 12700))
        w, h = int(r.width * 12700 * ratio), int(r.height * 12700 * ratio)
        slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), int((sw - w) / 2), int((sh - h) / 2), w, h)
        text = page.get_text("text", sort=True).strip()
        if text:
            slide.notes_slide.notes_text_frame.text = text
    prs.save(target)
    return target


def _export_via_libreoffice(pdf_bytes, fmt, target, pages, progress) -> str:
    if not externals.libreoffice_available():
        raise externals.ExternalToolMissing("Exporting to this format needs LibreOffice, which was not found.")
    tmpdir = tempfile.mkdtemp(prefix="pdfdesk-")
    try:
        docx_path = os.path.join(tmpdir, Path(target).stem + ".docx")
        _export_docx(pdf_bytes, docx_path, pages, None)
        _tick(progress, 1, 2, "Converting with LibreOffice")
        data = externals.libreoffice_convert(docx_path, fmt)
        with open(target, "wb") as fh:
            fh.write(data)
    finally:
        import shutil
        shutil.rmtree(tmpdir, ignore_errors=True)
    _tick(progress, 2, 2, "Done")
    return target


__all__ = ["ConversionError", "Cancelled", "to_pdf_bytes", "files_to_pdf_bytes", "images_to_pdf_bytes",
           "text_to_pdf_bytes", "html_to_pdf_bytes", "markdown_to_pdf_bytes", "export_pdf",
           "EXPORT_FORMATS", "import_filter", "can_import", "_unique_path"]
